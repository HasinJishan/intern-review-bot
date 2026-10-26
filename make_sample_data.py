"""Creates interns.xlsx (100 interns, 5 tasks) and 3 sample submissions for the demo.
Replace the sample names/emails in the Excel file with your real intern details."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

TASKS = [
    ("T1", "Palindrome Checker",
     "Write a Python function is_palindrome(text) that returns True if the text reads the same forwards and "
     "backwards, ignoring case, spaces and punctuation. Include at least 3 test cases.",
     '''import re

def is_palindrome(text: str) -> bool:
    cleaned = re.sub(r"[^a-z0-9]", "", text.lower())
    return cleaned == cleaned[::-1]

assert is_palindrome("A man, a plan, a canal: Panama")
assert not is_palindrome("hello")
assert is_palindrome("")
'''),
    ("T2", "Word Frequency Counter",
     "Write top_words(text, n) that returns the n most common words (lower-case, punctuation ignored) as a "
     "list of (word, count) tuples. Include test cases.",
     '''import re
from collections import Counter

def top_words(text: str, n: int = 3):
    words = re.findall(r"[a-z']+", text.lower())
    return Counter(words).most_common(n)

assert top_words("the cat and the hat. The end", 1) == [("the", 3)]
'''),
    ("T3", "Prime Numbers Generator",
     "Write primes_upto(n) that RETURNS a list of all prime numbers up to and including n. Handle n < 2. "
     "Include test cases.",
     '''def primes_upto(n: int) -> list[int]:
    if n < 2:
        return []
    sieve = [True] * (n + 1)
    sieve[0] = sieve[1] = False
    for i in range(2, int(n ** 0.5) + 1):
        if sieve[i]:
            for j in range(i * i, n + 1, i):
                sieve[j] = False
    return [i for i, is_prime in enumerate(sieve) if is_prime]

assert primes_upto(10) == [2, 3, 5, 7]
assert primes_upto(1) == []
'''),
    ("T4", "CSV Sales Totals",
     "Write totals_by_category(path) that reads a CSV with columns category,amount and returns a dict of "
     "total amount per category. Handle a missing file gracefully.",
     '''import csv
from collections import defaultdict

def totals_by_category(path):
    totals = defaultdict(float)
    try:
        with open(path, newline="") as f:
            for row in csv.DictReader(f):
                totals[row["category"]] += float(row["amount"])
    except FileNotFoundError:
        return {}
    return dict(totals)
'''),
    ("T5", "Bank Account Class",
     "Create a BankAccount class with deposit(amount), withdraw(amount) and a balance property. Withdrawing more "
     "than the balance must raise ValueError. Include test cases.",
     '''class BankAccount:
    def __init__(self, balance: float = 0):
        self._balance = balance

    @property
    def balance(self):
        return self._balance

    def deposit(self, amount):
        if amount <= 0:
            raise ValueError("Deposit must be positive")
        self._balance += amount

    def withdraw(self, amount):
        if amount > self._balance:
            raise ValueError("Insufficient funds")
        self._balance -= amount
'''),
]

FIRST = ["Aarav", "Diya", "Karthik", "Meera", "Rohan", "Sneha", "Vikram", "Ananya", "Arjun", "Priya"]
LAST = ["Kumar", "Sharma", "Reddy", "Iyer", "Nair", "Patel", "Singh", "Das", "Menon", "Rao"]

INTERN_COLS = [
    "ID", "Name", "Email", "Task_ID", "Status", "Kickoff_Sent", "Review_Date", "Review_Time", "Meet_Link",
    "Meet_ID", "Submission_File", "Score", "Feedback", "Improve_Areas", "Next_Review_Date",
    "Next_Review_Time", "Next_Meet_Link", "Next_Meet_ID",
]

wb = Workbook()
ws = wb.active
ws.title = "Interns"
ws.append(INTERN_COLS)
for i in range(100):
    first, last = FIRST[i % 10], LAST[i // 10]
    ws.append([f"I{i + 1:03d}", f"{first} {last}", f"{first}.{last}@example.com".lower(),
               TASKS[i % 5][0], "Not started"])

tasks = wb.create_sheet("Tasks")
tasks.append(["Task_ID", "Title", "Description", "Reference_Solution"])
for t in TASKS:
    tasks.append(list(t))

for sheet in (ws, tasks):
    for cell in sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
for col, w in zip("ABCDEFGHIJKLMNOPQR", [8, 20, 30, 9, 14, 18, 12, 14, 36, 14, 14, 8, 50, 40, 14, 14, 36, 14]):
    ws.column_dimensions[col].width = w
for col, w in zip("ABCD", [9, 26, 60, 70]):
    tasks.column_dimensions[col].width = w

# Legend so the sheet is self-explanatory
ws["T1"], ws["T2"] = "How to use", "Edit ID/Name/Email/Task_ID only. Columns E-R are filled by the bot."
ws["T1"].font = Font(name="Arial", bold=True)

wb.save("interns.xlsx")

# Three demo submissions: good (I001/T1), partial (I002/T2), weak (I003/T3)
Path("submissions").mkdir(exist_ok=True)
Path("submissions/I001.py").write_text('''import re

def is_palindrome(text):
    s = re.sub(r"[^a-z0-9]", "", text.lower())
    return s == s[::-1]

assert is_palindrome("Racecar")
assert is_palindrome("A man, a plan, a canal: Panama")
assert not is_palindrome("python")
''')
Path("submissions/I002.py").write_text('''def top_words(text, n):
    counts = {}
    for w in text.split(" "):
        if w in counts:
            counts[w] += 1
        else:
            counts[w] = 1
    return sorted(counts.items())
''')
Path("submissions/I003.py").write_text('''def primes_upto(n):
    for i in range(1, n):
        for j in range(2, i):
            if i % j == 0:
                break
        else:
            print(i)
''')
print("Created interns.xlsx (100 interns, 5 tasks) and submissions/I001-I003.py")
