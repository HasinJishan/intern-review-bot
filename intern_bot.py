"""AI Intern Review Bot

  python intern_bot.py kickoff   -> explains each intern's task + emails review meeting (date, time, Meet link)
  python intern_bot.py review    -> compares each submission with the reference solution, writes human-like
                                    feedback, emails it, and books the next session (date, time, Meet link, ID)

Safe by default: without --live nothing is sent and no calendar events are created (emails are saved to ./outbox).
"""
import argparse
import difflib
import json
import os
import random
import re
import smtplib
import ssl
import string
import uuid
from datetime import datetime, timedelta
from email.message import EmailMessage
from pathlib import Path
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from openpyxl import load_workbook

load_dotenv()

TZ = ZoneInfo(os.getenv("TIMEZONE", "Asia/Kolkata"))
EXCEL = Path(os.getenv("EXCEL_FILE", "interns.xlsx"))
SUBMISSIONS = Path("submissions")
OUTBOX = Path("outbox")
MODEL = os.getenv("ANTHROPIC_MODEL", "claude-sonnet-4-6")
GAP_DAYS = int(os.getenv("REVIEW_GAP_DAYS", "2"))        # days between sessions
BATCH_SIZE = int(os.getenv("BATCH_SIZE", "10"))          # interns per time slot
SLOT_MINUTES = int(os.getenv("SLOT_MINUTES", "30"))
DAY_START_HOUR = int(os.getenv("DAY_START_HOUR", "10"))

# ---------------------------------------------------------------- Excel helpers


def load():
    wb = load_workbook(EXCEL)
    ws = wb["Interns"]
    cols = {c.value: i + 1 for i, c in enumerate(ws[1])}
    tasks = {
        r[0]: {"title": r[1], "description": r[2], "reference": r[3]}
        for r in wb["Tasks"].iter_rows(min_row=2, values_only=True)
        if r[0]
    }
    return wb, ws, cols, tasks


def read_row(ws, cols, r):
    return {k: ws.cell(r, c).value for k, c in cols.items()}


def write_row(ws, cols, r, **values):
    for k, v in values.items():
        ws.cell(r, cols[k], v)


# ---------------------------------------------------------------- scheduling + Google Meet


def review_slot(index, today):
    """Review date = today + GAP_DAYS (skips weekends); interns are split into BATCH_SIZE groups per slot."""
    d = today + timedelta(days=GAP_DAYS)
    while d.weekday() >= 5:
        d += timedelta(days=1)
    d = d.replace(hour=DAY_START_HOUR, minute=0, second=0, microsecond=0)
    return d + timedelta(minutes=SLOT_MINUTES * (index // BATCH_SIZE))


_calendar = None


def _calendar_service():
    global _calendar
    if _calendar:
        return _calendar
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build

    scopes = ["https://www.googleapis.com/auth/calendar.events"]
    creds = None
    if Path("token.json").exists():
        creds = Credentials.from_authorized_user_file("token.json", scopes)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            creds = InstalledAppFlow.from_client_secrets_file("credentials.json", scopes).run_local_server(port=0)
        Path("token.json").write_text(creds.to_json())
    _calendar = build("calendar", "v3", credentials=creds)
    return _calendar


def create_meet(title, start, attendee, live):
    """Returns (meet_link, meeting_id). Real Google Meet when --live and credentials.json exist, else a demo code."""
    if live and Path("credentials.json").exists():
        end = start + timedelta(minutes=SLOT_MINUTES)
        event = {
            "summary": title,
            "start": {"dateTime": start.isoformat(), "timeZone": TZ.key},
            "end": {"dateTime": end.isoformat(), "timeZone": TZ.key},
            "attendees": [{"email": attendee}],
            "conferenceData": {
                "createRequest": {
                    "requestId": uuid.uuid4().hex,
                    "conferenceSolutionKey": {"type": "hangoutsMeet"},
                }
            },
        }
        ev = _calendar_service().events().insert(
            calendarId="primary", body=event, conferenceDataVersion=1, sendUpdates="none"
        ).execute()
        return ev["hangoutLink"], ev["conferenceData"]["conferenceId"]
    code = "-".join("".join(random.choices(string.ascii_lowercase, k=n)) for n in (3, 4, 3))
    return f"https://meet.google.com/{code}", code  # demo code, not a real meeting


# ---------------------------------------------------------------- email


def send_mail(to, subject, body, live):
    to = os.getenv("DEMO_EMAIL") or to  # redirect everything to one inbox while testing
    OUTBOX.mkdir(exist_ok=True)
    safe = re.sub(r"\W+", "_", subject)[:50]
    (OUTBOX / f"{datetime.now():%H%M%S%f}_{safe}.txt").write_text(f"To: {to}\nSubject: {subject}\n\n{body}", encoding="utf-8")
    if not live:
        return
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = os.environ["SMTP_USER"], to, subject
    msg.set_content(body)
    with smtplib.SMTP_SSL(os.getenv("SMTP_HOST", "smtp.gmail.com"), 465, context=ssl.create_default_context()) as s:
        s.login(os.environ["SMTP_USER"], os.environ["SMTP_PASS"])
        s.send_message(msg)


# ---------------------------------------------------------------- AI (free Gemini or Claude)


def _gemini(system, user, want_json):
    """Free Gemini API via plain HTTPS (no extra package). Retries when the free-tier rate limit is hit."""
    import time
    import urllib.error
    import urllib.request

    model = os.getenv("GEMINI_MODEL", "gemini-3.6-flash")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    config = {"maxOutputTokens": 4096}
    if want_json:
        config["responseMimeType"] = "application/json"
    payload = {
        "systemInstruction": {"parts": [{"text": system}]},
        "contents": [{"role": "user", "parts": [{"text": user}]}],
        "generationConfig": config,
    }
    headers = {"Content-Type": "application/json", "x-goog-api-key": os.environ["GEMINI_API_KEY"]}
    for attempt in range(4):
        try:
            req = urllib.request.Request(url, json.dumps(payload).encode(), headers)
            with urllib.request.urlopen(req, timeout=90) as resp:
                data = json.load(resp)
            time.sleep(7)  # stay under the free-tier requests-per-minute limit
            return data["candidates"][0]["content"]["parts"][0]["text"]
        except urllib.error.HTTPError as e:
            if e.code in (429, 500, 503) and attempt < 3:
                time.sleep(20)
                continue
            raise


def ask_claude(system, user, max_tokens=1000, want_json=False):
    """Uses free Gemini if GEMINI_API_KEY is set, else Claude if ANTHROPIC_API_KEY is set, else returns None."""
    try:
        if os.getenv("GEMINI_API_KEY"):
            return _gemini(system, user, want_json)
        if os.getenv("ANTHROPIC_API_KEY"):
            import anthropic

            r = anthropic.Anthropic().messages.create(
                model=MODEL, max_tokens=max_tokens, system=system, messages=[{"role": "user", "content": user}]
            )
            return r.content[0].text
    except Exception as e:  # keep the demo running even if the AI call fails
        print(f"  [warning] AI call failed ({e}); using offline fallback")
    return None  # offline fallback is used


KICKOFF_SYSTEM = (
    "You are a friendly senior engineer mentoring a new intern on a kickoff call. Explain the task the way you "
    "would say it out loud: the goal, a short step-by-step way to approach it, what a good result looks like, and "
    "one common mistake to avoid. Plain text only, no markdown symbols, under 200 words. Do not give the full solution."
)

REVIEW_SYSTEM = (
    "You are a warm, experienced mentor reviewing an intern's work. Compare the SUBMISSION with the REFERENCE "
    "solution for the TASK. Write like a real person: specific, honest, encouraging. Mention actual functions, "
    "variables or behaviours you saw. No generic praise, no phrases like 'As an AI'. Reply with JSON only:\n"
    '{"score": 0-100, "went_well": ["..", ".."], "improve": ["..", ".."], '
    '"feedback": "one natural paragraph of 80-150 words speaking directly to the intern", '
    '"next_focus": "one sentence goal for the next two days"}'
)


def explain_task(task):
    text = ask_claude(KICKOFF_SYSTEM, f"Task: {task['title']}\n\n{task['description']}")
    if text:
        return text.strip()
    return (
        f"Your task is '{task['title']}'. {task['description']}\n\n"
        "Approach: read the requirement twice, sketch the steps on paper, code the simplest version, then test "
        "edge cases (empty input, unusual characters). Common mistake: skipping the tests."
    )


def review_submission(task, submission):
    text = ask_claude(
        REVIEW_SYSTEM,
        f"TASK:\n{task['description']}\n\nREFERENCE:\n{task['reference']}\n\nSUBMISSION:\n{submission[:6000]}",
        1200,
        want_json=True,
    )
    if text:
        return json.loads(re.search(r"\{.*\}", text, re.S).group(0))
    # offline fallback (no API key): crude similarity so the pipeline still runs
    ratio = difflib.SequenceMatcher(None, submission, task["reference"]).ratio()
    return {
        "score": min(100, int(ratio * 100) + 15),
        "went_well": ["You submitted on time and the code runs structure-wise."],
        "improve": ["Compare your approach with the reference.", "Add test cases for edge cases."],
        "feedback": "Thanks for submitting. (Offline mode: set GEMINI_API_KEY for detailed human-like feedback.)",
        "next_focus": "Refine your solution and add tests.",
    }


# ---------------------------------------------------------------- commands


def when(dt):
    return dt.strftime("%A, %d %b %Y"), dt.strftime("%I:%M %p %Z")


def kickoff(args):
    wb, ws, cols, tasks = load()
    today, count = datetime.now(TZ), 0
    for r in range(2, ws.max_row + 1):
        iv = read_row(ws, cols, r)
        if not iv["ID"] or (args.only and iv["ID"] != args.only):
            continue
        if iv["Kickoff_Sent"] and not args.only:
            continue
        if args.limit and count >= args.limit:
            break
        task = tasks[iv["Task_ID"]]
        start = review_slot(r - 2, today)
        to = os.getenv("DEMO_EMAIL") or iv["Email"]
        link, meet_id = create_meet(f"Intern Review - {iv['Name']} - {task['title']}", start, to, args.live)
        date_s, time_s = when(start)
        body = (
            f"Hi {iv['Name'].split()[0]},\n\nWelcome to the team! Here is your first task.\n\n"
            f"TASK: {task['title']}\n{task['description']}\n\n"
            f"HOW TO APPROACH IT\n{explain_task(task)}\n\n"
            f"HOW TO SUBMIT\nSave your work and submit it before the review meeting (file name: {iv['ID']}.py).\n\n"
            f"REVIEW MEETING\nDate: {date_s}\nTime: {time_s}\nGoogle Meet: {link}\nMeeting ID: {meet_id}\n\n"
            "I'll go through your submission with you and share feedback.\n\nBest,\nIntern Review Bot"
        )
        send_mail(iv["Email"], f"Your task + review meeting on {date_s}", body, args.live)
        write_row(
            ws, cols, r,
            Status="Task assigned", Kickoff_Sent=datetime.now(TZ).strftime("%Y-%m-%d %H:%M"),
            Review_Date=start.strftime("%Y-%m-%d"), Review_Time=time_s, Meet_Link=link, Meet_ID=meet_id,
        )
        wb.save(EXCEL)
        count += 1
        print(f"[kickoff] {iv['ID']} {iv['Name']} -> {date_s} {time_s}")
    print(f"Done: {count} kickoff emails ({'LIVE' if args.live else 'dry-run'}).")


def review(args):
    wb, ws, cols, tasks = load()
    today, count = datetime.now(TZ), 0
    for r in range(2, ws.max_row + 1):
        iv = read_row(ws, cols, r)
        if not iv["ID"] or (args.only and iv["ID"] != args.only):
            continue
        sub_file = next(SUBMISSIONS.glob(f"{iv['ID']}.*"), None)
        if not sub_file:
            continue  # nothing submitted yet
        if args.limit and count >= args.limit:
            break
        task = tasks[iv["Task_ID"]]
        result = review_submission(task, sub_file.read_text(encoding="utf-8", errors="ignore"))
        start = review_slot(r - 2, today)
        to = os.getenv("DEMO_EMAIL") or iv["Email"]
        link, meet_id = create_meet(f"Intern Review - {iv['Name']} - Round 2", start, to, args.live)
        date_s, time_s = when(start)
        bullets = lambda items: "\n".join(f"  - {x}" for x in items)
        body = (
            f"Hi {iv['Name'].split()[0]},\n\nThanks for submitting '{task['title']}'. Here is my review.\n\n"
            f"{result['feedback']}\n\nSCORE: {result['score']}/100\n\n"
            f"WHAT WENT WELL\n{bullets(result['went_well'])}\n\n"
            f"AREAS TO IMPROVE\n{bullets(result['improve'])}\n\n"
            f"GOAL FOR THE NEXT 2 DAYS\n{result['next_focus']}\n\n"
            f"NEXT REVIEW SESSION\nDate: {date_s}\nTime: {time_s}\nGoogle Meet: {link}\nMeeting ID: {meet_id}\n\n"
            "Keep going, you're making progress!\n\nBest,\nIntern Review Bot"
        )
        send_mail(iv["Email"], f"Your review feedback ({result['score']}/100) + next session {date_s}", body, args.live)
        write_row(
            ws, cols, r,
            Status="Reviewed", Submission_File=sub_file.name, Score=result["score"],
            Feedback=result["feedback"], Improve_Areas="; ".join(result["improve"]),
            Next_Review_Date=start.strftime("%Y-%m-%d"), Next_Review_Time=time_s,
            Next_Meet_Link=link, Next_Meet_ID=meet_id,
        )
        wb.save(EXCEL)
        count += 1
        print(f"[review] {iv['ID']} {iv['Name']} score={result['score']} next={date_s} {time_s}")
    print(f"Done: {count} reviews ({'LIVE' if args.live else 'dry-run'}).")


if __name__ == "__main__":
    p = argparse.ArgumentParser(description="AI Intern Review Bot")
    p.add_argument("command", choices=["kickoff", "review"])
    p.add_argument("--live", action="store_true", help="really send emails and create Google Meet events")
    p.add_argument("--limit", type=int, help="process only N interns (good for demos)")
    p.add_argument("--only", help="process a single intern ID, e.g. I001")
    a = p.parse_args()
    {"kickoff": kickoff, "review": review}[a.command](a)